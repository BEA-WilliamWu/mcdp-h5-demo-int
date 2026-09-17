"""Compare SQL templates with DOCX v12; only HTML/layout whitespace is normalized."""
from pathlib import Path
import html
import re

ROOT = Path(__file__).resolve().parents[3]
SQL = ROOT / 'consulting/db/branch_change_history/20260917_BCOH2H-1288/1288_HTH_Management_Notification.sql'
DOCX = ROOT / 'story/notification/1.Customer Onboarding BM - Notification (Clear)-v12-20260917_192646.docx'


def values(sql, variable, append=False):
    literal = r"(q'\[.*?\]'|'(?:''|[^'])*')"
    prefix = re.escape(variable) + r'\s*:=\s*'
    if append:
        prefix += re.escape(variable) + r'\s*\|\|\s*'
    found = re.findall(prefix + literal + r'\s*;', sql, re.S)
    return [v[3:-2] if v.startswith("q'[") else v[1:-1].replace("''", "'") for v in found]


def plain(value):
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', value))).strip()


def templates():
    sql = SQL.read_text()
    subjects = values(sql, 'v_subject')
    bodies = values(sql, 'v_body')
    sms = {lang: values(sql, 'v_sms_' + lang) for lang in ('en', 'tc', 'sc')}
    assert len(subjects) == len(bodies) == 3
    assert all(len(items) == 3 for items in sms.values())
    return {action: {'subject': subjects[i], 'email': bodies[i], 'sms': sms['en'][i],
                     'sms_locales': {lang: items[i] for lang, items in sms.items()}}
            for i, action in enumerate(('ENABLE', 'DISABLE', 'EDIT'))}


def verify():
    from docx import Document
    doc = Document(DOCX)
    for i, (action, item) in enumerate(templates().items(), 1):
        email = doc.tables[1].rows[i].cells
        sms = doc.tables[2].rows[i].cells
        assert plain(item['subject']) == plain(email[2].text), action + ': bilingual subject differs'
        assert plain(item['email']) == plain(email[3].text), action + ': bilingual body differs'
        assert len(item['subject'].encode('utf-8')) <= 200
        assert len(item['email'].encode('utf-8')) <= 4000
        for lang, column in (('en', 2), ('tc', 3), ('sc', 4)):
            value = item['sms_locales'][lang]
            assert value == sms[column].text, action + '/' + lang + ': SMS differs'
            assert len(value.encode('utf-8')) <= 1000
            print(f"PASS: {action}/{lang} exact DOCX SMS; length={len(value)}")
        print(f"PASS: {action} complete bilingual subject/body match DOCX v12")
    sql = SQL.read_text()
    assert "'zh-hant' THEN v_sms_tc" in sql and "'zh-hans-cn' THEN v_sms_sc" in sql
    assert "d.destination || '_' || l.locale" in sql
    assert "'M', 'CZ', l.locale" in sql and 'v_count <> 18' in sql


if __name__ == '__main__':
    verify()
